"""Standard-only: in-app basic DOCX editor.

Loading uses the same Mammoth render as annotation, so what the user edits is
what they saw. Saving rebuilds a DOCX from the rich-text document (paragraphs,
headings, lists, bold/italic/underline, images, basic tables). Headers,
footers, comments, revisions and complex sections are not round-tripped; the
window shows that limitation and the default save target is a copy.
"""
from __future__ import annotations

import os
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import mammoth
from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches, Pt, RGBColor
from PySide6.QtCore import QUrl, Qt, Signal
from PySide6.QtGui import (QAction, QFont, QImage, QKeySequence, QTextCharFormat,
                           QTextCursor, QTextDocument, QTextImageFormat,
                           QTextListFormat)
from PySide6.QtWidgets import (QComboBox, QFileDialog, QLabel, QMainWindow,
                               QMessageBox, QTextEdit, QToolBar, QVBoxLayout,
                               QWidget)

from sketchpad.ui import theme

EDITOR_CSS = """
body { font-family: 'Microsoft YaHei', 'Segoe UI', sans-serif;
       font-size: 11pt; color: #111111; line-height: 1.5; }
h1 { font-size: 20pt; margin: 12pt 0 6pt 0; }
h2 { font-size: 16pt; margin: 10pt 0 5pt 0; }
h3 { font-size: 13pt; margin: 8pt 0 4pt 0; }
p { margin: 4pt 0; }
pre, code { font-family: Consolas, monospace; background: #f2f2f2; }
table { border-collapse: collapse; }
td, th { border: 1px solid #bbbbbb; padding: 4pt; }
th { background: #f0f0f0; }
img { margin: 4pt 0; }
"""

FIDELITY_HINT = "建议「另存为」。"
IMAGE_MAX_WIDTH = 480
# Mammoth ignores underline/strike by default; keep them so the editor matches Word.
STYLE_MAP = "u => u\nstrike => s"


class WordEditError(ValueError):
    """Editing failed with a message meant for the status bar/dialog."""


@dataclass
class SourceInfo:
    path: Path
    fidelity_reasons: list[str] = field(default_factory=list)
    layout: dict = field(default_factory=dict)


def _read_layout(path: Path) -> dict:
    """直接从 document.xml 读页面设置：不打开 python-docx，避免锁住源文件。"""
    try:
        with zipfile.ZipFile(path) as archive:
            body = archive.read("word/document.xml").decode("utf-8", "ignore")
    except (OSError, KeyError, zipfile.BadZipFile):
        return {}
    size = re.search(r"<w:pgSz\b([^>]*)/?>", body)
    if size is None:
        return {}
    attrs = dict(re.findall(r'w:(\w+)="([^"]*)"', size.group(1)))
    try:
        width = Emu(int(attrs.get("w", 11906)) * 635)
        height = Emu(int(attrs.get("h", 16838)) * 635)
    except (TypeError, ValueError):
        return {}
    margins = {}
    margin_match = re.search(r"<w:pgMar\b([^>]*)/?>", body)
    if margin_match is not None:
        values = dict(re.findall(r'w:(\w+)="([^"]*)"', margin_match.group(1)))
        for name, key in (("top_margin", "top"), ("bottom_margin", "bottom"),
                          ("left_margin", "left"), ("right_margin", "right")):
            try:
                margins[name] = Emu(int(values[key]) * 635)
            except (KeyError, TypeError, ValueError):
                continue
    landscape = attrs.get("orient") == "landscape"
    return {"page_width": width, "page_height": height,
            "orientation": WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT,
            "margins": margins}


def _fidelity_reasons(path: Path) -> list[str]:
    reasons: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            body = ""
            if "word/document.xml" in names:
                body = archive.read("word/document.xml").decode("utf-8", "ignore")
    except (OSError, zipfile.BadZipFile):
        return ["文件结构无法完整解析"]
    if any(name.startswith("word/header") for name in names):
        reasons.append("页眉")
    if any(name.startswith("word/footer") for name in names):
        reasons.append("页脚")
    if "word/comments.xml" in names:
        reasons.append("批注")
    if "<w:ins " in body or "<w:del " in body:
        reasons.append("修订")
    if "word/footnotes.xml" in names or "word/endnotes.xml" in names:
        reasons.append("脚注/尾注")
    if "<w:txbxContent" in body:
        reasons.append("文本框")
    if body.count("<w:sectPr") > 1:
        reasons.append("多节分页")
    if "<w:gridSpan" in body or "<w:vMerge" in body:
        reasons.append("合并单元格")
    if "<m:oMath" in body:
        reasons.append("公式")
    return reasons


def load_document(path, assets_dir) -> tuple[QTextDocument, SourceInfo]:
    path = Path(path)
    if path.suffix.lower() != ".docx":
        raise WordEditError("只能编辑 .docx；旧版 .doc 请先另存为 .docx")
    if not path.is_file():
        raise WordEditError(f"找不到 Word 文件：{path}")
    assets = Path(assets_dir)
    assets.mkdir(parents=True, exist_ok=True)
    counter = {"value": 0}

    def save_image(image) -> dict:
        counter["value"] += 1
        content_type = (image.content_type or "image/png").split("/")[-1]
        extension = "jpg" if content_type in ("jpeg", "jpg") else content_type
        target = assets / f"image-{counter['value']}.{extension}"
        with image.open() as handle, open(target, "wb") as out:
            out.write(handle.read())
        return {"src": target.as_uri()}

    try:
        # mammoth 内部打开 docx 后不关 zip；传文件对象并由 with 关闭，
        # 否则 Windows 上原 Word 文件会被锁住，覆盖保存时报 WinError 5。
        with open(path, "rb") as handle:
            result = mammoth.convert_to_html(
                handle, convert_image=mammoth.images.img_element(save_image),
                style_map=STYLE_MAP)
    except Exception as exc:
        raise WordEditError(f"无法读取 Word 文件：{exc}") from exc
    document = QTextDocument()
    document.setDefaultFont(QFont("Microsoft YaHei", 11))
    document.setDocumentMargin(24)
    document.setHtml(f"<html><head><style>{EDITOR_CSS}</style></head>"
                     f"<body>{result.value}</body></html>")
    document.setModified(False)
    return document, SourceInfo(path, _fidelity_reasons(path), _read_layout(path))


def _iter_children(frame):
    iterator = frame.begin()
    while not iterator.atEnd():
        block = iterator.currentBlock()
        if block.isValid():
            yield "block", block
        else:
            yield "frame", iterator.currentFrame()
        iterator += 1


def _local_path(name: str) -> Path | None:
    if not name:
        return None
    url = QUrl(name)
    if url.isLocalFile():
        return Path(url.toLocalFile())
    candidate = Path(name)
    return candidate if candidate.exists() else None


def _merge_alignment(paragraph, alignment) -> None:
    if alignment & Qt.AlignHCenter:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif alignment & Qt.AlignRight:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif alignment & Qt.AlignJustify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _list_style(block, docx_doc) -> str | None:
    text_list = block.textList()
    if text_list is None:
        return None
    ordered = text_list.format().style() in (
        QTextListFormat.ListDecimal, QTextListFormat.ListLowerAlpha,
        QTextListFormat.ListUpperAlpha, QTextListFormat.ListLowerRoman,
        QTextListFormat.ListUpperRoman)
    name = "List Number" if ordered else "List Bullet"
    level = max(0, int(block.blockFormat().indent()))
    if level:
        name = f"{name} {min(level + 1, 3)}"
    return name if name in docx_doc.styles else ("List Number" if ordered else "List Bullet")


def _write_paragraph(paragraph, block, docx_doc) -> None:
    heading = int(block.blockFormat().headingLevel())
    if heading:
        style = f"Heading {min(heading, 9)}"
        if style in docx_doc.styles:
            paragraph.style = style
    else:
        style = _list_style(block, docx_doc)
        if style and style in docx_doc.styles:
            paragraph.style = style
    _merge_alignment(paragraph, block.blockFormat().alignment())

    iterator = block.begin()
    while not iterator.atEnd():
        fragment = iterator.fragment()
        if not fragment.isValid():
            iterator += 1
            continue
        char_format = fragment.charFormat()
        text = fragment.text()
        if char_format.isImageFormat():
            image_format = char_format.toImageFormat()
            image_path = _local_path(image_format.name())
            if image_path is not None and image_path.is_file():
                run = paragraph.add_run()
                width = float(image_format.width() or 0.0)
                if width > 0:
                    run.add_picture(str(image_path), width=Pt(width * 0.75))
                else:
                    run.add_picture(str(image_path), width=Inches(5.5))
            iterator += 1
            continue
        if text:
            run = paragraph.add_run(text.replace("\u2028", "\n"))
            run.bold = char_format.fontWeight() >= 600
            run.italic = char_format.fontItalic()
            run.underline = char_format.fontUnderline()
            if char_format.fontPointSize() > 0:
                run.font.size = Pt(char_format.fontPointSize())
            families = char_format.fontFamilies()
            if families:
                run.font.name = families[0]
            brush = char_format.foreground()
            if brush.style() != Qt.NoBrush:
                red, green, blue, _alpha = brush.color().getRgb()
                run.font.color.rgb = RGBColor(red, green, blue)
        iterator += 1


def _write_table(table_frame, docx_doc) -> None:
    rows, columns = table_frame.rows(), table_frame.columns()
    table = docx_doc.add_table(rows=rows, cols=columns)
    try:
        table.style = "Table Grid"
    except (KeyError, ValueError):
        pass
    for row in range(rows):
        for column in range(columns):
            cell = table.cell(row, column)
            first = True
            for kind, item in _iter_children(table_frame.cellAt(row, column)):
                if kind == "block":
                    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    _write_paragraph(paragraph, item, docx_doc)


def _write_container(frame, docx_doc) -> None:
    for kind, item in _iter_children(frame):
        if kind == "block":
            _write_paragraph(docx_doc.add_paragraph(), item, docx_doc)
        elif hasattr(item, "rows") and hasattr(item, "columns"):
            _write_table(item, docx_doc)
        elif item is not None:
            _write_container(item, docx_doc)


def _apply_layout(new_doc, layout: dict | None) -> None:
    if not layout:
        return
    try:
        section = new_doc.sections[0]
        if layout.get("page_width") is not None:
            section.page_width = layout["page_width"]
        if layout.get("page_height") is not None:
            section.page_height = layout["page_height"]
        if layout.get("orientation") is not None:
            section.orientation = layout["orientation"]
        for name, value in (layout.get("margins") or {}).items():
            if value is not None:
                setattr(section, name, value)
    except Exception:
        pass


def save_document(text_document: QTextDocument, target, source=None,
                  layout: dict | None = None) -> Path:
    """把编辑器内容写成 .docx（先写临时文件再替换，避免写坏目标）。"""
    target = Path(target)
    if target.suffix.lower() != ".docx":
        raise WordEditError("只能保存为 .docx 文件")
    docx_doc = DocxDocument()
    if layout is None and source is not None:
        layout = _read_layout(Path(source))
    _apply_layout(docx_doc, layout)
    _write_container(text_document.rootFrame(), docx_doc)
    if not docx_doc.paragraphs and not docx_doc.tables:
        docx_doc.add_paragraph()
    partial = target.with_name(target.name + ".tmp")
    try:
        docx_doc.save(str(partial))
        os.replace(partial, target)
    except Exception as exc:
        partial.unlink(missing_ok=True)
        raise WordEditError(f"保存 Word 失败：{exc}") from exc
    return target


class WordEditorWindow(QMainWindow):
    """Non-modal basic DOCX editor; default save target is a copy."""

    saved = Signal(str)

    def __init__(self, path, parent=None) -> None:
        super().__init__(parent)
        self.source_path = Path(path)
        self.target_path: Path | None = None
        self._dirty = False
        self._loading = True
        self._temp_dir = Path(tempfile.mkdtemp(prefix="sketchpad-word-edit-"))
        self.document, self.source_info = load_document(
            self.source_path, self._temp_dir / "assets")
        self.resize(920, 760)
        self._build_ui()
        self._refresh_title()
        self._loading = False

    # --- UI ---

    def _build_ui(self) -> None:
        root = QWidget(self)
        layout = QVBoxLayout(root)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        self.banner = QLabel("", root)
        self.banner.setWordWrap(True)
        self.banner.setStyleSheet(
            f"background: #4a3b12; color: #ffd60a; padding: 6px 10px;")
        if self.source_info.fidelity_reasons:
            self.banner.setText(
                "基础编辑：编辑/保存不会保留 "
                + "、".join(self.source_info.fidelity_reasons)
                + f"；{FIDELITY_HINT}")
        self.banner.setVisible(bool(self.source_info.fidelity_reasons))
        layout.addWidget(self.banner)

        self.editor = QTextEdit(root)
        self.editor.setAcceptRichText(True)
        self.editor.setDocument(self.document)
        self.editor.setStyleSheet(
            "QTextEdit { background: #ffffff; color: #111111;"
            " border: none; padding: 24px 48px; }")

        toolbar = QToolBar(root)
        toolbar.setMovable(False)
        layout.addWidget(toolbar)
        self._build_actions(toolbar)
        layout.addWidget(self.editor, 1)
        self.setCentralWidget(root)
        self.editor.textChanged.connect(self._on_text_changed)
        self.editor.cursorPositionChanged.connect(self._sync_heading_combo)

    def _action(self, toolbar, text, slot, shortcut=None, checkable=False) -> QAction:
        action = QAction(text, self)
        action.setCheckable(checkable)
        tip = text.rstrip("…")
        if shortcut:
            tip = f"{tip}（{shortcut}）"
        action.setToolTip(tip)
        action.setStatusTip(tip)
        if shortcut:
            action.setShortcut(QKeySequence(shortcut))
        action.triggered.connect(slot)
        toolbar.addAction(action)
        return action

    def _build_actions(self, toolbar) -> None:
        self.undo_action = self._action(toolbar, "撤销", self.editor.undo, "Ctrl+Z")
        self.redo_action = self._action(toolbar, "重做", self.editor.redo, "Ctrl+Y")
        toolbar.addSeparator()
        self.bold_action = self._action(toolbar, "加粗", self._toggle_bold, "Ctrl+B", True)
        self.italic_action = self._action(toolbar, "斜体", self._toggle_italic, "Ctrl+I", True)
        self.underline_action = self._action(
            toolbar, "下划线", self._toggle_underline, "Ctrl+U", True)
        toolbar.addSeparator()
        self.heading_combo = QComboBox(toolbar)
        for label in ("正文", "标题 1", "标题 2", "标题 3"):
            self.heading_combo.addItem(label)
        self.heading_combo.activated.connect(self._apply_heading)
        toolbar.addWidget(self.heading_combo)
        toolbar.addSeparator()
        self._action(toolbar, "项目符号", lambda: self._apply_list(ordered=False))
        self._action(toolbar, "编号", lambda: self._apply_list(ordered=True))
        self._action(toolbar, "插入图片…", self._insert_image)
        toolbar.addSeparator()
        self._action(toolbar, "保存", self.save, "Ctrl+S")
        self._action(toolbar, "另存为…", self.save_as, "Ctrl+Shift+S")

    # --- formatting ---

    def _merge_format(self, char_format: QTextCharFormat) -> None:
        cursor = self.editor.textCursor()
        if cursor.hasSelection():
            cursor.mergeCharFormat(char_format)
            self.editor.mergeCurrentCharFormat(char_format)
        else:
            self.editor.mergeCurrentCharFormat(char_format)

    def _toggle_bold(self) -> None:
        fmt = QTextCharFormat()
        bold = self.editor.currentCharFormat().fontWeight() >= 600
        fmt.setFontWeight(QFont.Normal if bold else QFont.Bold)
        self._merge_format(fmt)

    def _toggle_italic(self) -> None:
        fmt = QTextCharFormat()
        fmt.setFontItalic(not self.editor.currentCharFormat().fontItalic())
        self._merge_format(fmt)

    def _toggle_underline(self) -> None:
        fmt = QTextCharFormat()
        fmt.setFontUnderline(not self.editor.currentCharFormat().fontUnderline())
        self._merge_format(fmt)

    def _apply_heading(self, index: int) -> None:
        level = index
        cursor = self.editor.textCursor()
        cursor.beginEditBlock()
        try:
            start, end = cursor.selectionStart(), cursor.selectionEnd()
            probe = QTextCursor(self.document)
            probe.setPosition(start)
            while True:
                block_format = probe.blockFormat()
                block_format.setHeadingLevel(level)
                probe.setBlockFormat(block_format)
                if probe.block().position() + probe.block().length() >= end:
                    break
                if not probe.movePosition(QTextCursor.NextBlock):
                    break
            char_format = QTextCharFormat()
            if level:
                char_format.setFontWeight(QFont.Bold)
                char_format.setFontPointSize({1: 20, 2: 16, 3: 13}.get(level, 12))
            else:
                char_format.setFontWeight(QFont.Normal)
                char_format.setFontPointSize(11)
            if cursor.hasSelection():
                cursor.mergeCharFormat(char_format)
            else:
                self.editor.mergeCurrentCharFormat(char_format)
        finally:
            cursor.endEditBlock()
        self.editor.setFocus()

    def _apply_list(self, ordered: bool) -> None:
        cursor = self.editor.textCursor()
        current = cursor.currentList()
        desired = QTextListFormat.ListDecimal if ordered else QTextListFormat.ListDisc
        cursor.beginEditBlock()
        try:
            if current is not None and current.format().style() == desired:
                current.remove(cursor.block())       # 再点一次同类型：取消列表
            elif current is not None:
                list_format = current.format()
                list_format.setStyle(desired)        # 项目符号 ⇄ 编号：切换整组的样式
                current.setFormat(list_format)
            else:
                list_format = QTextListFormat()
                list_format.setStyle(desired)
                cursor.createList(list_format)
        finally:
            cursor.endEditBlock()
        self.editor.setFocus()

    def _insert_image(self) -> None:
        chosen, _filter = QFileDialog.getOpenFileName(
            self, "插入图片", str(self.source_path.parent),
            "图片 (*.png *.jpg *.jpeg *.bmp *.gif)")
        if not chosen:
            return
        source = Path(chosen)
        images = self._temp_dir / "inserted"
        images.mkdir(parents=True, exist_ok=True)
        target = images / f"image-{len(list(images.iterdir())) + 1}{source.suffix}"
        try:
            shutil.copy2(source, target)
        except OSError as exc:
            self._warn(f"图片复制失败：{exc}")
            return
        image = QImage(str(target))
        image_format = QTextImageFormat()
        image_format.setName(str(target))
        if image.isNull():
            image_format.setWidth(IMAGE_MAX_WIDTH)
        else:
            width = min(image.width(), IMAGE_MAX_WIDTH)
            image_format.setWidth(width)
            image_format.setHeight(width * image.height() / max(1, image.width()))
        self.editor.textCursor().insertImage(image_format)
        self.editor.setFocus()

    def _sync_heading_combo(self) -> None:
        level = int(self.editor.textCursor().blockFormat().headingLevel())
        self.heading_combo.setCurrentIndex(min(level, 3))

    # --- dirty / title ---

    def _on_text_changed(self) -> None:
        if self._loading:
            return
        self._dirty = True
        self._refresh_title()

    def _refresh_title(self) -> None:
        mark = "● " if self._dirty else ""
        self.setWindowTitle(f"{mark}编辑 — {self.source_path.name}")

    # --- save ---

    def _default_target(self) -> Path:
        if self.target_path is not None:
            return self.target_path
        return self.source_path.with_name(f"{self.source_path.stem}-编辑.docx")

    def _confirm_overwrite(self, target: Path) -> bool:
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Warning)
        box.setWindowTitle("覆盖原 Word 文件")
        box.setText(f"要用编辑后的内容覆盖 {target.name} 吗？")
        box.setInformativeText(
            "会先把原文件备份为 "
            f"{target.name}.bak。页眉页脚等复杂结构可能不会保留。")
        box.setStandardButtons(QMessageBox.Yes | QMessageBox.No)
        box.setDefaultButton(QMessageBox.No)
        return box.exec() == QMessageBox.Yes

    def _choose_target(self) -> Path | None:
        chosen, _filter = QFileDialog.getSaveFileName(
            self, "保存 Word", str(self._default_target()), "Word 文档 (*.docx)")
        if not chosen:
            return None
        target = Path(chosen)
        if target.suffix.lower() != ".docx":
            target = target.with_suffix(".docx")
        return target

    def save(self) -> bool:
        """首次保存询问副本路径；之后 Ctrl+S 直接写同一目标。"""
        target = self.target_path or self._choose_target()
        if target is None:
            return False
        return self._write(target)

    def save_as(self) -> bool:
        target = self._choose_target()
        return bool(target) and self._write(target)

    def _write(self, target: Path) -> bool:
        target = Path(target)
        if target.resolve() == self.source_path.resolve():
            if not self._confirm_overwrite(target):
                return False
            try:
                shutil.copy2(self.source_path, self.source_path.with_name(
                    self.source_path.name + ".bak"))
            except OSError as exc:
                self._warn(f"备份原文件失败：{exc}")
                return False
        try:
            save_document(self.document, target, self.source_path,
                          layout=self.source_info.layout)
        except WordEditError as exc:
            self._warn(str(exc))
            return False
        self.target_path = target
        self._dirty = False
        self._refresh_title()
        self.saved.emit(str(target))
        return True

    # --- lifecycle ---

    def _ask_save_changes(self) -> int:
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Warning)
        box.setWindowTitle("未保存的修改")
        box.setText("Word 编辑内容尚未保存。")
        box.setStandardButtons(QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel)
        box.setDefaultButton(QMessageBox.Save)
        return box.exec()

    def closeEvent(self, event) -> None:
        if self._dirty:
            choice = self._ask_save_changes()
            if choice == QMessageBox.Cancel:
                event.ignore()
                return
            if choice == QMessageBox.Save and not self.save():
                event.ignore()
                return
        shutil.rmtree(self._temp_dir, ignore_errors=True)
        super().closeEvent(event)

    def _warn(self, text: str) -> None:
        QMessageBox.warning(self, "Word 编辑", text)
