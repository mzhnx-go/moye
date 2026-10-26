"""Word 基础编辑：加载、复杂结构提示、往返保存、覆盖备份。"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import hashlib
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image
from PySide6.QtGui import QTextCursor
from PySide6.QtWidgets import QApplication

from sketchpad.word import editor as word_editor
from sketchpad.word.editor import WordEditorWindow, load_document, save_document

_ALIVE: list = []


@pytest.fixture(scope="module")
def qapp():
    yield QApplication.instance() or QApplication([])


def _make_docx(path, *, header=False, extra=False):
    doc = Document()
    if header:
        doc.sections[0].header.paragraphs[0].text = "页眉文字"
    doc.add_heading("第一章 标题", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("正文")
    paragraph.add_run("粗体").bold = True
    paragraph.add_run("斜体").italic = True
    paragraph.add_run("下划线").underline = True
    doc.add_paragraph("项目一", style="List Bullet")
    doc.add_paragraph("项目二", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "表头A"
    table.cell(1, 1).text = "值2"
    png = path.with_name("dot.png")
    Image.new("RGB", (32, 24), (10, 120, 200)).save(png)
    doc.add_picture(str(png), width=Inches(0.5))
    if extra:
        doc.add_paragraph("备注")
    doc.save(str(path))
    return path


def test_load_document_roundtrips_rich_content(qapp, tmp_path):
    source = _make_docx(tmp_path / "讲义.docx")
    document, info = load_document(source, tmp_path / "assets")
    assert info.fidelity_reasons == []
    assert "第一章" in document.toPlainText()
    assert "粗体" in document.toPlainText()

    target = tmp_path / "副本.docx"
    save_document(document, target, source)
    result = Document(str(target))
    texts = [p.text for p in result.paragraphs]
    assert "第一章 标题" in texts
    assert result.paragraphs[0].style.name.startswith("Heading")
    assert any(run.bold for p in result.paragraphs for run in p.runs if run.text == "粗体")
    assert any(run.italic for p in result.paragraphs for run in p.runs if run.text == "斜体")
    assert any(run.underline for p in result.paragraphs for run in p.runs if run.text == "下划线")
    assert any(p.style.name.startswith("List") for p in result.paragraphs)
    assert result.tables[0].cell(1, 1).text == "值2"
    assert len(result.inline_shapes) == 1
    assert source.read_bytes()  # 原文件仍在


def test_editor_edit_and_save_copy_keeps_source(qapp, tmp_path):
    source = _make_docx(tmp_path / "合同.docx")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    window = WordEditorWindow(source)
    _ALIVE.append(window)
    cursor = window.document.find("正文")
    assert not cursor.isNull()
    cursor.insertText("改后的正文")
    target = tmp_path / "合同-编辑.docx"
    window._choose_target = lambda: target
    assert window.save() is True
    assert window.target_path == target
    assert target.is_file()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert "改后的正文" in "\n".join(p.text for p in Document(str(target)).paragraphs)
    window.close()


def test_save_over_original_creates_backup(qapp, tmp_path):
    source = _make_docx(tmp_path / "覆盖.docx")
    original = source.read_bytes()
    window = WordEditorWindow(source)
    _ALIVE.append(window)
    problems = []
    window._warn = problems.append
    window.document.find("正文").insertText("新内容")
    window._choose_target = lambda: source
    window._confirm_overwrite = lambda target: True
    assert window.save() is True, problems
    backup = source.with_name(source.name + ".bak")
    assert backup.read_bytes() == original
    texts = [p.text for p in Document(str(source)).paragraphs]
    assert "新内容" in "\n".join(texts)
    window.close()


def test_fidelity_banner_lists_unsupported_parts(qapp, tmp_path):
    source = _make_docx(tmp_path / "复杂.docx", header=True)
    window = WordEditorWindow(source)
    _ALIVE.append(window)
    assert "页眉" in window.source_info.fidelity_reasons
    assert "页眉" in window.banner.text()
    window._dirty = False
    window.close()


def test_toolbar_actions_change_saved_docx(qapp, tmp_path, monkeypatch):
    source = _make_docx(tmp_path / "工具栏.docx")
    window = WordEditorWindow(source)
    _ALIVE.append(window)
    found = window.document.find("正文")
    cursor = QTextCursor(window.document)
    cursor.setPosition(found.selectionStart())
    cursor.setPosition(found.selectionStart() + 2, QTextCursor.KeepAnchor)
    window.editor.setTextCursor(cursor)
    window._toggle_bold()
    window._toggle_italic()
    window._toggle_underline()

    heading_cursor = window.document.find("项目一")
    window.editor.setTextCursor(heading_cursor)
    window._apply_heading(1)

    number_cursor = window.document.find("项目二")
    window.editor.setTextCursor(number_cursor)
    window._apply_list(ordered=True)
    end = QTextCursor(window.document)
    end.movePosition(QTextCursor.End)
    window.editor.setTextCursor(end)

    png = tmp_path / "new.png"
    Image.new("RGB", (20, 20), (0, 200, 0)).save(png)
    monkeypatch.setattr(
        word_editor.QFileDialog, "getOpenFileName",
        staticmethod(lambda *args, **kwargs: (str(png), "图片")))
    window._insert_image()

    target = tmp_path / "工具栏-编辑.docx"
    window._choose_target = lambda: target
    assert window.save() is True
    result = Document(str(target))
    changed = [run for p in result.paragraphs for run in p.runs
               if run.text.startswith("正文")]
    assert changed and changed[0].bold and changed[0].italic and changed[0].underline
    assert any(p.text == "项目一" and p.style.name == "Heading 1"
               for p in result.paragraphs)
    assert any(p.text == "项目二" and p.style.name.startswith("List Number")
               for p in result.paragraphs), [
                   (p.text, p.style.name) for p in result.paragraphs]
    assert len(result.inline_shapes) == 2
    window.close()


def test_default_save_target_is_a_copy(qapp, tmp_path):
    source = _make_docx(tmp_path / "默认.docx")
    window = WordEditorWindow(source)
    _ALIVE.append(window)
    assert window._default_target().name == "默认-编辑.docx"
    window._dirty = False
    window.close()
