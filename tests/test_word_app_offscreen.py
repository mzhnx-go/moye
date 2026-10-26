"""Word 批注接线：缓存 PDF、sidecar 绑原 docx、查找/导出、旧 .doc 报错。"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
import pymupdf
from docx import Document
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

import main
from sketchpad.core import pdfstore
from sketchpad.core.strokes import Stroke
from sketchpad.reader import install as install_readers
from sketchpad.word import install as install_word
from sketchpad.word import converter
from sketchpad.word.view import WordView
from sketchpad.workspace.viewers import clear_viewers

_ALIVE: list = []


@pytest.fixture(scope="module")
def qapp():
    yield QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def viewers(tmp_path, monkeypatch):
    install_readers()
    install_word()
    monkeypatch.setattr(converter, "cache_base", lambda: tmp_path / "converted")
    yield
    clear_viewers()


@pytest.fixture
def docx(tmp_path):
    path = tmp_path / "合同.docx"
    doc = Document()
    doc.add_heading("第一章 条款", level=1)
    doc.add_paragraph("数量关系与资料分析 ABC")
    doc.add_paragraph("第二段内容")
    doc.save(str(path))
    return path


@pytest.fixture
def pdf(tmp_path):
    path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    doc.new_page(width=595, height=842).insert_text((40, 60), "PDF")
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def win(qapp, tmp_path):
    ctrl = main.AppController(tmp_path / "cfg")
    w = main.MainWindow(ctrl)
    w._confirm_unsaved = lambda action: QMessageBox.Discard
    w.show()
    QTest.qWaitForWindowExposed(w)
    _ALIVE.append(w)
    yield w
    w.scratch.hide()
    w._quitting = True
    w.close()


def test_docx_opens_as_cached_pdf(win, docx):
    assert win._enter_document(docx) is True
    view = win.pdf_view
    assert isinstance(view, WordView)
    assert view.kind == "pdf"
    assert view.pdf_path == docx                # 标题/最近/sidecar 都指向原文件
    assert view.search_path != docx
    assert view.search_path.suffix.lower() == ".pdf"
    assert view.search_path.is_file()
    assert docx.read_bytes()                    # 原 Word 不变
    assert "缓存 PDF" in win.message_label.text()
    assert win.tabs.current_entry.title == docx.name


def test_word_edit_entry_points_follow_current_document(win, docx, pdf):
    win._enter_document(docx)
    assert win.menubar.edit_word_action.isEnabled() is True
    assert win.toolbox.word_button.isHidden() is False
    assert win._edit_word_shortcut.isEnabled() is True
    win._enter_document(pdf)
    assert win.menubar.edit_word_action.isEnabled() is False
    assert win.toolbox.word_button.isHidden() is True
    assert win._edit_word_shortcut.isEnabled() is False
    win._enter_draft()
    assert win.menubar.edit_word_action.isEnabled() is False
    assert win.toolbox.word_button.isHidden() is True
    assert win._edit_word_shortcut.isEnabled() is False


def test_edit_word_opens_in_app_editor(win, docx, tmp_path):
    win._enter_document(docx)
    assert win._edit_word_document() is True
    window = next(iter(win._word_editors.values()))
    assert window.source_path == docx
    cursor = window.document.find("数量关系")
    assert not cursor.isNull()
    cursor.insertText("改后的内容")
    target = tmp_path / "合同-编辑.docx"
    window._choose_target = lambda: target
    assert window.save() is True
    assert target.is_file()
    assert "改后的内容" in window.document.toPlainText()
    window.close()


def test_sidecar_is_written_next_to_docx(win, docx):
    win._enter_document(docx)
    page = win.pdf_view.document.pages[0]
    page.strokes.append(Stroke(points=[(10.0, 10.0), (60.0, 40.0)],
                               color="#e8eaed", width=2.0))
    assert win._save_annotations() is True
    sidecar = pdfstore.sidecar_path(docx)
    assert sidecar.is_file()
    assert sidecar.parent == docx.parent
    assert docx.with_name(docx.name + ".sketch.source").is_file()
    assert list(win.pdf_view.search_path.parent.glob("*.sketch.json")) == []


def test_changed_word_backs_up_stale_annotations(win, docx):
    win._enter_document(docx)
    page = win.pdf_view.document.pages[0]
    page.strokes.append(Stroke(points=[(10.0, 10.0), (60.0, 40.0)],
                               color="#e8eaed", width=2.0))
    assert win._save_annotations() is True
    win._close_current_entry()

    from docx import Document as NewDocument
    changed = NewDocument()
    changed.add_paragraph("改过之后的内容")
    changed.save(str(docx))

    assert win._enter_document(docx) is True
    older = list(docx.parent.glob("*.sketch.json.stale-*"))
    assert older
    assert "内容已修改" in win.message_label.text()
    assert win.pdf_view.document.pages[0].strokes == []


def test_search_and_highlight_run_on_cached_pdf(win, docx):
    win._enter_document(docx)
    win._find_search("数量关系")
    assert win._find_hits
    assert all(page == 0 for page, _ in win._find_hits)


def test_export_uses_original_stem(win, docx, tmp_path, monkeypatch):
    win._enter_document(docx)
    seen = {}

    def fake_save(parent, caption, default, filt):
        seen["default"] = default
        out = tmp_path / "导出.pdf"
        return (str(out), filt)

    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(fake_save))
    assert win._export_pdf() is True
    assert seen["default"].endswith("合同-批注.pdf")
    assert (tmp_path / "导出.pdf").is_file()


def test_reopen_reuses_converted_pdf(win, docx):
    win._enter_document(docx)
    first = win.pdf_view.search_path
    stamp = first.stat().st_mtime_ns
    win._close_current_entry()
    assert win._enter_document(docx) is True
    assert win.pdf_view.search_path == first
    assert first.stat().st_mtime_ns == stamp


def test_legacy_doc_reports_actionable_error(win, tmp_path):
    legacy = tmp_path / "旧稿.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    assert win._enter_document(legacy) is False
    assert win.mode == "draft"
    assert "另存为 .docx" in win.message_label.text()
