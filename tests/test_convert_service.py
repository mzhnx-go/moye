"""转换中心服务层：合并/拆分/旋转/压缩/图片/文本/DOCX/Markdown/页码/水印/加密。"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from PIL import Image

from sketchpad.convert import service


def _pdf(path, pages=2):
    document = pymupdf.open()
    for index in range(pages):
        page = document.new_page(width=595, height=842)
        page.insert_text((40, 60), f"Page {index + 1} content")
    document.save(str(path))
    document.close()
    return path


def test_merge_and_split(tmp_path):
    first = _pdf(tmp_path / "a.pdf", 2)
    second = _pdf(tmp_path / "b.pdf", 3)
    merged = service.merge_pdfs([first, second], tmp_path / "merged.pdf")
    with pymupdf.open(merged) as doc:
        assert doc.page_count == 5

    parts = service.split_pdf(merged, tmp_path / "parts", ranges="1-2,5")
    assert len(parts) == 2
    with pymupdf.open(parts[0]) as doc:
        assert doc.page_count == 2
    with pymupdf.open(parts[1]) as doc:
        assert doc.page_count == 1
    assert first.is_file() and second.is_file()


def test_rotate_and_compress(tmp_path):
    source = _pdf(tmp_path / "rotate.pdf", 1)
    rotated = service.rotate_pdf(source, tmp_path / "rotated.pdf", angle=90)
    with pymupdf.open(rotated) as doc:
        assert doc[0].rotation == 90
    compressed = service.compress_pdf(source, tmp_path / "compressed.pdf")
    with pymupdf.open(compressed) as doc:
        assert doc.page_count == 1
        assert "Page 1" in doc[0].get_text()


def test_pdf_images_text_and_docx(tmp_path):
    source = _pdf(tmp_path / "doc.pdf", 2)
    images = service.pdf_to_images(source, tmp_path / "img", dpi=96, fmt="png")
    assert len(images) == 2 and all(item.stat().st_size > 0 for item in images)

    text = service.pdf_to_text(source, tmp_path / "doc.txt")
    assert "Page 2" in text.read_text(encoding="utf-8")

    docx = service.pdf_to_docx(source, tmp_path / "doc.docx")
    paragraphs = [p.text for p in DocxDocument(str(docx)).paragraphs]
    assert any("Page 1" in text for text in paragraphs)


def test_images_to_pdf(tmp_path):
    first = tmp_path / "one.png"
    second = tmp_path / "two.jpg"
    Image.new("RGB", (120, 90), (255, 0, 0)).save(first)
    Image.new("RGB", (90, 120), (0, 0, 255)).save(second)
    output = service.images_to_pdf([first, second], tmp_path / "images.pdf")
    with pymupdf.open(output) as doc:
        assert doc.page_count == 2


def test_markdown_and_docx_to_pdf(tmp_path):
    markdown = tmp_path / "note.md"
    markdown.write_text("# 标题\n\n正文 **粗体**\n\n- 一\n- 二\n", encoding="utf-8")
    output = service.markdown_to_pdf(markdown, tmp_path / "note.pdf")
    with pymupdf.open(output) as doc:
        assert "标题" in doc[0].get_text()

    docx = tmp_path / "word.docx"
    document = DocxDocument()
    document.add_heading("Word 标题", level=1)
    document.add_paragraph("Word 正文")
    document.save(str(docx))
    pdf = service.docx_to_pdf(docx, tmp_path / "word.pdf")
    with pymupdf.open(pdf) as doc:
        assert "Word 正文" in doc[0].get_text()


def test_page_numbers_and_watermark(tmp_path):
    source = _pdf(tmp_path / "marks.pdf", 2)
    numbered = service.add_page_numbers(
        source, tmp_path / "numbered.pdf", position="bottom-center")
    with pymupdf.open(numbered) as doc:
        text = "".join(page.get_text() for page in doc)
        assert "1" in text and "2" in text
    marked = service.add_watermark(source, tmp_path / "marked.pdf", text="内部")
    with pymupdf.open(marked) as doc:
        assert "内部" in "".join(page.get_text() for page in doc)


def test_encrypt_and_decrypt(tmp_path):
    source = _pdf(tmp_path / "secret.pdf", 1)
    encrypted = service.encrypt_pdf(
        source, tmp_path / "locked.pdf", user_password="open", owner_password="own")
    with pymupdf.open(encrypted) as doc:
        assert bool(doc.needs_pass) is True
        assert doc.authenticate("wrong") == 0
    decrypted = service.decrypt_pdf(
        encrypted, tmp_path / "unlocked.pdf", password="open")
    with pymupdf.open(decrypted) as doc:
        assert bool(doc.needs_pass) is False
        assert "Page 1" in doc[0].get_text()


def test_ocr_adds_text_layer(tmp_path):
    from PIL import Image as PILImage, ImageDraw

    image_path = tmp_path / "scan.png"
    image = PILImage.new("RGB", (600, 200), "white")
    ImageDraw.Draw(image).text((30, 80), "Hello 123", fill="black")
    image.save(image_path)
    scan = pymupdf.open()
    page = scan.new_page(width=600, height=200)
    page.insert_image(page.rect, filename=str(image_path))
    scan.save(str(tmp_path / "scan.pdf"))
    scan.close()

    output = service.ocr_pdf(tmp_path / "scan.pdf", tmp_path / "ocr.pdf")
    with pymupdf.open(output) as document:
        text = document[0].get_text()
    assert "123" in text
