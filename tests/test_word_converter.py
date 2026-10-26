"""Word → 缓存 PDF：内容/图片落盘、指纹复用、旧 .doc 明确报错。"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image
from PySide6.QtWidgets import QApplication

from sketchpad.word import converter


@pytest.fixture(scope="module")
def qapp():
    yield QApplication.instance() or QApplication([])


def _make_docx(path, text="正文内容", image: bool = False):
    doc = Document()
    doc.add_heading("第一章 标题", level=1)
    doc.add_paragraph(text)
    if image:
        png = path.with_name("dot.png")
        Image.new("RGB", (40, 40), (200, 30, 30)).save(png)
        doc.add_picture(str(png), width=Inches(0.6))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表头A"
    table.cell(0, 1).text = "表头B"
    table.cell(1, 0).text = "值1"
    table.cell(1, 1).text = "值2"
    doc.save(str(path))
    return path


def test_docx_converts_to_pdf_with_text(qapp, tmp_path):
    source = _make_docx(tmp_path / "讲义.docx", "数量关系与资料分析")
    pdf = converter.ensure_cached_pdf(source, tmp_path / "cache")
    assert pdf.is_file() and pdf.stat().st_size > 0
    assert source.read_bytes()  # 原文件仍在
    with pymupdf.open(pdf) as doc:
        text = "".join(page.get_text() for page in doc)
    assert "第一章" in text
    assert "数量关系" in text
    assert "表头A" in text


def test_embedded_image_is_extracted(qapp, tmp_path):
    source = _make_docx(tmp_path / "图.docx", image=True)
    pdf = converter.ensure_cached_pdf(source, tmp_path / "cache")
    assets = list(pdf.parent.joinpath("assets").glob("image-*"))
    assert assets and assets[0].stat().st_size > 0
    with pymupdf.open(pdf) as doc:
        assert any(page.get_images() for page in doc)


def test_cache_is_keyed_by_content_hash(qapp, tmp_path):
    source = _make_docx(tmp_path / "笔记.docx", "第一版")
    base = tmp_path / "cache"
    first = converter.ensure_cached_pdf(source, base)
    stamp = first.stat().st_mtime_ns
    again = converter.ensure_cached_pdf(source, base)
    assert again == first
    assert again.stat().st_mtime_ns == stamp

    _make_docx(source, "第二版")
    second = converter.ensure_cached_pdf(source, base)
    assert second != first
    assert second.is_file()


def test_legacy_doc_gets_actionable_error(qapp, tmp_path):
    source = tmp_path / "旧稿.doc"
    source.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(converter.WordConversionError) as info:
        converter.ensure_cached_pdf(source, tmp_path / "cache")
    assert "另存为 .docx" in str(info.value)
